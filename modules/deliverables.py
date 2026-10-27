"""Two separate outputs: complete research XLSX and provisional editorial PDF."""
from io import BytesIO
from html import escape
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import cm

NAVY='#10103F'; CYAN='#16C6DD'; PURPLE='#7D3CE6'; PINK='#EF258E'
DATA_KEYS=['keywords','clasificacion','facetas','sources','rejected','entities','concepts','wikidata','openalex','serp','competidores']

def excel_report(run):
    """Genera XLSX en memoria, compatible con Streamlit Community Cloud."""
    wb = Workbook()
    sh = wb.active
    sh.title = 'Resumen ejecutivo'
    sh.merge_cells('A1:H2')
    sh['A1'] = 'INVESTIGACIÓN SEMÁNTICA | ' + str(run['topic'])
    sh['A1'].font = Font(name='Aptos Display', size=17, bold=True, color='FFFFFF')
    sh['A1'].alignment = Alignment(vertical='center')
    for row in sh['A1:H2']:
        for cell in row: cell.fill = PatternFill('solid', fgColor=NAVY[1:])
    sh.merge_cells('A4:H4')
    sh['A4'] = str(run['country']) + ' | Español | Demanda CSV Semrush'
    sh['A4'].font = Font(bold=True, color=NAVY[1:])
    sh['A4'].fill = PatternFill('solid', fgColor='E8F8FC')
    measures = [('Keywords','keywords'),('Facetas provisionales','facetas'),('Fuentes','sources'),('Entidades','entities'),('Conceptos','concepts'),('Wikidata candidatos','wikidata'),('OpenAlex candidatos','openalex')]
    for i,(label,key) in enumerate(measures):
        row = 6 + i*2
        sh.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4)
        sh.merge_cells(start_row=row,start_column=5,end_row=row,end_column=8)
        sh.cell(row,1,label)
        sh.cell(row,5,len(run.get(key,[])))
        for col in range(1,9):
            sh.cell(row,col).fill = PatternFill('solid',fgColor='F1F0F8')
        sh.cell(row,1).font = Font(bold=True,color=NAVY[1:])
        sh.cell(row,5).font = Font(bold=True,size=14,color=PURPLE[1:])
    sh.merge_cells('A22:H24')
    sh['A22']='Facetas y QIDs exploratorios. El informe editorial no sustituye validación de intención, canibalización ni inventario de URLs.'
    sh['A22'].alignment=Alignment(wrap_text=True,vertical='center')
    sh['A22'].fill=PatternFill('solid',fgColor='FCEAF4')
    for col in range(1,9): sh.column_dimensions[get_column_letter(col)].width=16
    sh.sheet_view.showGridLines=False
    for key in DATA_KEYS:
        frame=run.get(key)
        if frame is None: continue
        ws=wb.create_sheet(key[:31]); ws.sheet_view.showGridLines=False
        columns=list(frame.columns) if len(frame.columns) else ['estado']
        ws.append(columns)
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor=NAVY[1:]);cell.font=Font(bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True)
        ws.row_dimensions[1].height=28
        for row in frame.itertuples(index=False,name=None):
            values=[]
            for v in row:
                if v is None or str(v).lower() in ('nan','nat','none'): values.append(None)
                elif isinstance(v,(int,float,bool)): values.append(v)
                else:
                    value=str(v)[:32000]
                    # Evitar que textos de CSV externos se interpreten como fórmulas.
                    if value.startswith(('=','+','-','@')): value="'"+value
                    values.append(value)
            ws.append(values)
        for idx,col in enumerate(columns,1):
            ws.column_dimensions[get_column_letter(idx)].width=min(48,max(16,len(str(col))+4))
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
    buffer=BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def editorial_plan(run,domain='',existing_urls=''):
    """Editorial hypotheses, not verified page decisions."""
    topic=run['topic'].strip(); country=run['country']
    domain=domain.strip().replace('https://','').replace('http://','').strip('/')
    existing=[x.strip() for x in existing_urls.splitlines() if x.strip()]
    facets=run.get('facetas')
    names=facets['grupo'].astype(str).tolist() if facets is not None and 'grupo' in facets else []
    names=[x for x in names if 'verificar' not in x.lower()][:6]
    if not names:names=['Conceptos fundamentales','Preguntas frecuentes']
    kw=run.get('keywords')
    top=kw.sort_values('volume',ascending=False).head(12) if kw is not None and 'volume' in kw else None
    return {'topic':topic,'country':country,'domain':domain,'existing':existing,'facets':names,'top':top,
      'status':'HIPOTESIS PARA REVISION: no se han validado intenciones ni URLs existentes',
      'note':'No crear ni actualizar URLs hasta contrastar inventario actual, intencion y riesgo de canibalizacion.'}

def editorial_pdf(run,domain='',existing_urls=''):
    p=editorial_plan(run,domain,existing_urls)
    out=BytesIO(); doc=SimpleDocTemplate(out,pagesize=(21*cm,29.7*cm),leftMargin=1.8*cm,rightMargin=1.8*cm,topMargin=1.8*cm,bottomMargin=1.6*cm)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='HeroSEO',parent=styles['Title'],fontName='Helvetica-Bold',fontSize=20,leading=25,textColor=colors.HexColor(NAVY),spaceAfter=14))
    styles.add(ParagraphStyle(name='HeadSEO',parent=styles['Heading2'],fontSize=12,leading=16,textColor=colors.HexColor(PURPLE),spaceBefore=15,spaceAfter=8))
    styles.add(ParagraphStyle(name='BodySEO',parent=styles['BodyText'],fontSize=9,leading=14,spaceAfter=8))
    styles.add(ParagraphStyle(name='SmallSEO',parent=styles['BodyText'],fontSize=8,leading=11))
    def P(t,style='BodySEO'):return Paragraph(escape(str(t)),styles[style])
    story=[P('INFORME EDITORIAL | BRIEF Y ARQUITECTURA','HeroSEO'),P(f"Tema: {p['topic']}   |   Mercado: {p['country']}"),P(p['status'])]
    story += [P('01  Resumen ejecutivo','HeadSEO'),P('La investigación aporta datos de demanda y conceptos semánticos. Las recomendaciones de contenido siguientes son hipótesis editoriales, no decisiones de publicación validadas.')]
    story += [P('02  Brief SEO preliminar','HeadSEO'),P('Objetivo: responder a la necesidad principal del tema y cubrir facetas respaldadas por el corpus y las keywords.'),P('Audiencia: usuarios que investigan el tema en el mercado seleccionado. Definir segmento específico antes de redactar.'),P('Keyword principal candidata: '+p['topic']+'. Validar intención y demanda exacta antes de asignarla.')]
    if p['top'] is not None:
        cells=[[P('Keyword','SmallSEO'),P('Volumen','SmallSEO'),P('Intención Semrush','SmallSEO')]]
        for _,r in p['top'].head(8).iterrows():cells.append([P(r.get('keyword',''),'SmallSEO'),P(r.get('volume',''),'SmallSEO'),P(r.get('intent','Sin dato'),'SmallSEO')])
        t=Table(cells,colWidths=[9*cm,2.2*cm,5.8*cm],repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(NAVY)),('TEXTCOLOR',(0,0),(-1,0),colors.white),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F3F8')]),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7)]));story.append(t)
    story += [P('03  Estructura H1–H3 propuesta','HeadSEO'),P('H1: '+p['topic']),P('H2: Qué es y cómo funciona'),P('H2: Aspectos esenciales para '+p['country'])]
    for f in p['facets'][:5]:story.append(P('H3 (hipótesis): '+f))
    story += [P('H2: Preguntas frecuentes (seleccionar preguntas reales de Semrush; no inventar demanda)')]
    story += [P('04  Arquitectura de URLs propuesta','HeadSEO')]
    slug=re.sub(r'[^a-z0-9]+','-',__import__('unicodedata').normalize('NFKD',p['topic'].lower()).encode('ascii','ignore').decode()).strip('-')
    suggested=('/'+slug+'/')
    story += [P('Página pilar candidata: '+(f'https://{p["domain"]}{suggested}' if p['domain'] else suggested)),P('Estado: URL sugerida, NO verificada como disponible ni necesaria.')]
    for f in p['facets'][:4]:
        sf=re.sub(r'[^a-z0-9]+','-',__import__('unicodedata').normalize('NFKD',f.lower()).encode('ascii','ignore').decode()).strip('-')
        story.append(P('Subtema candidato: '+suggested+sf+'/ — revisar si debe ser sección H2 o URL independiente.'))
    story += [P('URLs existentes aportadas: '+str(len(p['existing']))),P('No se ha realizado crawling del dominio ni validación de canibalización. El inventario proporcionado solo permite una revisión preliminar.')]
    story += [P('05  Enlazado y validación antes de publicar','HeadSEO'),P('Vincular páginas relacionadas solo cuando se compruebe que existen y sus intenciones son distintas.'),P(p['note'])]
    story += [P('06  Límites y trazabilidad','HeadSEO'),P('Wikidata: candidatos hasta validación de identidad. OpenAlex: metadatos no equivalen a lectura científica. Grupos: hipótesis, no clusters SERP. Volumen: estimación de demanda, no tráfico potencial.')]
    def footer(canvas,doc):
        canvas.setStrokeColor(colors.HexColor(CYAN));canvas.line(1.8*cm,1.2*cm,19.2*cm,1.2*cm)
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor(NAVY));canvas.drawString(1.8*cm,.9*cm,'INVESTIGACION SEMANTICA  |  INFORME EDITORIAL');canvas.drawRightString(19.2*cm,.9*cm,str(doc.page))
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()
