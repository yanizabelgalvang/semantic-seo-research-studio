import io
import pandas as pd
from modules.architecture import parse_sitemap, facet_matrix, plan_pages, brief_rows, verify_entities
from modules.deliverables import full_excel,editorial_pdf_v2

def sample():
    kw=pd.DataFrame([{'keyword':'posicionamiento web','volume':260,'kd':60,'intent':'Informational','cpc':1.2,'grupo':'Intención por verificar'},{'keyword':'servicio de posicionamiento seo','volume':1000,'kd':31,'intent':'Commercial','cpc':2,'grupo':'Contratación y servicios'},{'keyword':'agencia seo','volume':200,'kd':20,'intent':'Commercial','cpc':2,'grupo':'Contratación y servicios'},{'keyword':'consultor seo','volume':100,'kd':10,'intent':'Commercial','cpc':2,'grupo':'Contratación y servicios'}])
    q=pd.DataFrame([{'keyword':'como contratar una agencia seo','volume':10,'kd':None,'intent':'Informational','cpc':None}])
    return kw,q

def test_sitemap_and_proposals():
    sm=parse_sitemap(b'<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://ejemplo.com/servicio-de-posicionamiento-seo</loc></url></urlset>')
    k,q=sample();m=facet_matrix(k,q);assert m.preguntas_observadas.sum()==1
    p=plan_pages('posicionamiento web','México',k,q,sm,pd.DataFrame(),pd.DataFrame(),pd.DataFrame())
    assert len(p)==2
    assert 'AUDITAR' in p.iloc[1].accion
    b=brief_rows(p,k,q,pd.DataFrame(),pd.DataFrame(),pd.DataFrame())
    assert {'H1','Fuentes citadas','Firma del autor'}.issubset(set(b.bloque))
    run={'topic':'posicionamiento web','country':'México','keywords':k,'questions':q,'clasificacion':k,'facetas':m,'matriz':m,'sitemap':sm,'arquitectura':p,'brief':b}
    x=full_excel(run);pdf=editorial_pdf_v2(run)
    assert x[:2]==b'PK' and pdf[:4]==b'%PDF'
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(x),read_only=True)
    assert 'brief' in wb.sheetnames and 'arquitectura' in wb.sheetnames

def test_qid_requires_human_review():
    f=pd.DataFrame([{'qid_candidato':'Q180711','termino':'SEO'}])
    assert verify_entities(f,{}).iloc[0].estado_final=='pendiente'
    assert verify_entities(f,{'Q180711':'aprobar'}).iloc[0].estado_final=='verificado_por_usuario'
